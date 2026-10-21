from __future__ import annotations

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

SOURCE = Path('/dev/shm/cp337-final/CareSuite_Ambulant_A-Z_Arbeitsliste_2026-10-08.xlsx')
OUTPUT = Path('/dev/shm/cp338-work/CareSuite_Ambulant_A-Z_Arbeitsliste_2026-10-08.xlsx')
CATALOG = Path('docs/audit/ambulant-az-katalog-2026-10-08.json')
TASK_IDS = tuple(f'A-{number:02d}' for number in range(1, 9))
PACKET_IDS = ('A-P01', 'A-P02', 'A-P03')
ALLOWED_CELLS = {'Aufgaben': {f'{column}{row}' for row in range(9, 17) for column in ('M', 'O')}, 'Pakete': {f'{column}{row}' for row in range(8, 11) for column in ('G', 'H')}}

catalog = json.loads(CATALOG.read_text())
tasks = {row['id']: row for row in catalog['tasks'] if row['id'] in TASK_IDS}
packets = {row[0]: row for row in catalog['packets'] if row[0] in PACKET_IDS}
before = load_workbook(SOURCE, data_only=False)
after = load_workbook(OUTPUT, data_only=False)

assert before.sheetnames == after.sheetnames
changed_values = []
changed_heights = []
formula_count_before = formula_count_after = 0
for before_sheet, after_sheet in zip(before.worksheets, after.worksheets, strict=True):
    assert before_sheet.freeze_panes == after_sheet.freeze_panes
    assert before_sheet.sheet_view.showGridLines == after_sheet.sheet_view.showGridLines
    assert before_sheet.sheet_state == after_sheet.sheet_state
    for row in range(1, max(before_sheet.max_row, after_sheet.max_row) + 1):
        if before_sheet.row_dimensions[row].height != after_sheet.row_dimensions[row].height:
            changed_heights.append(f'{before_sheet.title}!{row}')
    for row in range(1, max(before_sheet.max_row, after_sheet.max_row) + 1):
        for column in range(1, max(before_sheet.max_column, after_sheet.max_column) + 1):
            a = before_sheet.cell(row, column)
            b = after_sheet.cell(row, column)
            if isinstance(a.value, str) and a.value.startswith('='):
                formula_count_before += 1
            if isinstance(b.value, str) and b.value.startswith('='):
                formula_count_after += 1
            assert a.style_id == b.style_id, f'Stil geändert: {before_sheet.title}!{a.coordinate}'
            if a.value != b.value:
                assert a.coordinate in ALLOWED_CELLS.get(before_sheet.title, set()), f'Unerlaubte Wertänderung: {before_sheet.title}!{a.coordinate}'
                changed_values.append(f'{before_sheet.title}!{a.coordinate}')

assert set(changed_values) == {f'{sheet}!{cell}' for sheet, cells in ALLOWED_CELLS.items() for cell in cells}
assert set(changed_heights) == {f'Aufgaben!{row}' for row in range(9, 17)} | {f'Pakete!{row}' for row in range(8, 11)}
assert formula_count_before == formula_count_after == 264

task_rows = {after['Aufgaben'].cell(row, 1).value: row for row in range(1, after['Aufgaben'].max_row + 1)}
for task_id in TASK_IDS:
    row = task_rows[task_id]
    assert after['Aufgaben'].cell(row, 13).value == tasks[task_id]['evidence']
    assert after['Aufgaben'].cell(row, 15).value == tasks[task_id]['next']
    assert after['Aufgaben'].row_dimensions[row].height == 300
packet_rows = {after['Pakete'].cell(row, 1).value: row for row in range(1, after['Pakete'].max_row + 1)}
for packet_id in PACKET_IDS:
    row = packet_rows[packet_id]
    assert after['Pakete'].cell(row, 7).value == packets[packet_id][6]
    assert after['Pakete'].cell(row, 8).value == packets[packet_id][7]
    assert after['Pakete'].row_dimensions[row].height == 300

with ZipFile(SOURCE) as first, ZipFile(OUTPUT) as second:
    assert first.namelist() == second.namelist()
    allowed_members = {'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'}
    for name in first.namelist():
        if name not in allowed_members:
            assert first.read(name) == second.read(name), f'ZIP-Mitglied geändert: {name}'

digest = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
print(json.dumps({'changedValues': len(changed_values), 'changedRowHeights': len(changed_heights), 'formulas': formula_count_after, 'unchangedZipMembers': len(ZipFile(OUTPUT).namelist()) - 2, 'bytes': OUTPUT.stat().st_size, 'sha256': digest}, ensure_ascii=False))
