from __future__ import annotations

import textwrap
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

BEFORE = Path('/dev/shm/cp337-final/CareSuite_Ambulant_A-Z_Arbeitsliste_2026-10-08.xlsx')
AFTER = Path('/dev/shm/cp338-work/CareSuite_Ambulant_A-Z_Arbeitsliste_2026-10-08.xlsx')
ROOT = Path('/dev/shm/cp338-rendered-views')
FONT_PATH = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
BOLD_PATH = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
FONT = ImageFont.truetype(FONT_PATH, 18)
SMALL = ImageFont.truetype(FONT_PATH, 16)
BOLD = ImageFont.truetype(BOLD_PATH, 21)


def wrapped(value: object, width: int) -> list[str]:
    text = '' if value is None else str(value)
    lines: list[str] = []
    for paragraph in text.splitlines() or ['']:
        lines.extend(textwrap.wrap(paragraph, width=width, break_long_words=False, replace_whitespace=False) or [''])
    return lines


def draw_view(workbook_path: Path, stage: str, sheet_name: str, row: int, columns: tuple[int, int], item_id: str) -> None:
    workbook = load_workbook(workbook_path, data_only=False)
    sheet = workbook[sheet_name]
    left = wrapped(sheet.cell(row, columns[0]).value, 92)
    right = wrapped(sheet.cell(row, columns[1]).value, 92)
    height = 150 + max(len(left), len(right)) * 25
    image = Image.new('RGB', (2200, height), '#F5F8FC')
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((24, 24, 2176, height - 24), radius=28, fill='#FFFFFF', outline='#93B4DC', width=3)
    draw.text((55, 48), f'{sheet_name} · {item_id} · {stage}', font=BOLD, fill='#103B68')
    draw.text((55, 88), 'Nachweis', font=BOLD, fill='#205D91')
    draw.text((1125, 88), 'Nächster Schritt', font=BOLD, fill='#205D91')
    draw.line((1098, 86, 1098, height - 50), fill='#C8D9EB', width=2)
    y_left = y_right = 130
    for line in left:
        draw.text((55, y_left), line, font=SMALL, fill='#182635')
        y_left += 25
    for line in right:
        draw.text((1125, y_right), line, font=SMALL, fill='#182635')
        y_right += 25
    target = ROOT / stage
    target.mkdir(parents=True, exist_ok=True)
    image.save(target / f'{item_id}.png')


for stage, path in (('before', BEFORE), ('after', AFTER)):
    for row in range(9, 17):
        draw_view(path, stage, 'Aufgaben', row, (13, 15), f'A-{row - 8:02d}')
    for row in range(8, 11):
        draw_view(path, stage, 'Pakete', row, (7, 8), f'A-P0{row - 7}')

print(f'22 Ansichten gerendert: {ROOT}')
